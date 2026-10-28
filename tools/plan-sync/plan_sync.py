#!/usr/bin/env python3
"""Plan workbook -> text mirror and tracker (Jira) import files.   (see docs/process/PLAN_TRACKER_WORKFLOW.md)

    pip install openpyxl
    python tools/plan-sync/plan_sync.py export                 # plan/backlog.csv + plan/jira-import.csv + plan/jira-links.csv
    python tools/plan-sync/plan_sync.py check                  # offline consistency checks on the workbook (exit 1 on findings)

The workbook itself is tracker-free; everything tracker-specific (field names, sprint import, assignees) lives in this tool.
`export` is read-only on the workbook. `push`/`pull`/tracker-side `check` need API access and are planned (backlog #336); their
credentials are CI secrets (JIRA_BASE_URL, JIRA_PROJECT_KEY, JIRA_USER_EMAIL, JIRA_API_TOKEN), never files in git.
Assignees come from plan/members.json ({"Member A": "a@company.com", ...}); copy plan/members.example.json. The real file is git-ignored.
"""
import csv, json, re, sys
from collections import Counter
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "plan" / "PEMS_Backlog_Roadmap.xlsx"
OUT_DIR = ROOT / "plan"
MEMBERS = OUT_DIR / "members.json"

ISSUE_TYPE = {"Feature": "Story", "Bug Fix": "Bug"}                    # everything else -> Task
PRIORITY = {"P0": "Highest", "P1": "High", "P2": "Medium", "P3": "Low"}
STATUS = {"Not Started": "To Do", "In Progress": "In Progress", "Partial": "In Progress", "Done": "Done",
          "Hold": "On Hold", "Later": "Backlog", "Cancelled": "Won't Do", "Obsolete": "Won't Do"}
EPICS = {"M0": "M0 Baseline and foundation decisions", "M1": "M1 Dynamic core, components and licensing (Alpha)",
         "M2": "M2 Builders, demo and install (Beta)", "M3": "M3 Commercial release 1.0",
         "M4": "M4 Platform completion and module migration", "M5": "M5 Competitive features",
         "M6": "M6 Later", "-": "No milestone"}
LABEL_COLUMNS = 6


def load_rows():
    wb = openpyxl.load_workbook(WORKBOOK, data_only=False)
    ws = wb["Backlog"]
    header = [c.value for c in ws[1]]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if values[0] is None:
            continue
        rows.append(dict(zip(header, values)))
    return header, rows


def clean(v):
    return "" if v is None else str(v).strip()


def load_members():
    if MEMBERS.exists():
        return json.loads(MEMBERS.read_text(encoding="utf-8"))
    return {}


def sprint_of(r):
    s = clean(r.get("Sprint"))
    return "" if not s or "stretch" in s.lower() else s


def labels(r):
    out = []
    if clean(r.get("Effort")): out.append("effort-" + clean(r["Effort"]))
    if clean(r.get("Doer")): out.append("doer-" + clean(r["Doer"]).lower())
    if clean(r.get("Confirmed by PM")).lower() == "yes": out.append("pm-confirmed")
    if clean(r.get("Status")) == "Partial": out.append("partial")
    if clean(r.get("Status")) == "Later": out.append("later")
    if clean(r.get("Commercial")) == "Sales blocker": out.append("sales-blocker")
    if "stretch" in clean(r.get("Sprint")).lower(): out.append("sprint-1-stretch")
    # team-managed Jira spaces have no Components: the plan area travels as a label instead
    area = re.sub(r"[^a-z0-9]+", "-", clean(r.get("Area")).lower()).strip("-")
    if area: out.append("area-" + area[:60])
    return out


def summary(r):
    return f"[#{r['#']}] {clean(r.get('Task'))}"


def description(r, titles):
    head = (f"*Plan ID:* #{r['#']}   *Area:* {clean(r.get('Area'))}   *Group:* {clean(r.get('Group'))}   "
            f"*Milestone:* {clean(r.get('Milestone')) or '-'}   *Priority:* {clean(r.get('Priority'))}   "
            f"*Effort:* {clean(r.get('Effort'))}   *Owner:* {clean(r.get('Owner')) or '-'}   *Doer:* {clean(r.get('Doer'))}")
    parts = [head, "*Step 0:* read AGENTS.md and docs/guidelines/README.md, then the related Jira issues, before starting."]
    for title, key in (("Why", "Evidence"), ("What to do", "Prompt"), ("Done when", "Requirement")):
        if clean(r.get(key)):
            text = clean(r[key])
            if key == "Requirement" and text.lower().startswith("done when:"):
                text = text[len("done when:"):].strip()          # the section header already says "Done when"
            parts.append(f"h3. {title}\n{text}")
    deps = re.findall(r"#(\d+)", clean(r.get("Depends")))
    if deps:
        parts.append("h3. Depends on\n" + "\n".join(f"* [#{d}] {titles.get(int(d), '(unknown)')}" for d in deps))
    if clean(r.get("Report")):
        parts.append(f"h3. Report\n{clean(r['Report'])}")
    return "\n\n".join(parts)


def export():
    header, rows = load_rows()
    OUT_DIR.mkdir(exist_ok=True)
    members = load_members()
    titles = {int(r["#"]): clean(r.get("Task")) for r in rows}

    with (OUT_DIR / "backlog.csv").open("w", newline="", encoding="utf-8") as fh:        # reviewable text mirror
        w = csv.writer(fh)
        w.writerow(header)
        for r in rows:
            w.writerow([clean(r.get(h)) for h in header])

    milestones = sorted({clean(r.get("Milestone")) or "-" for r in rows})
    with (OUT_DIR / "jira-import.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["Work item ID", "Work type", "Parent", "Summary", "Description", "Priority", "Status", "Fix Version", "Plan ID",
                    "Sprint", "Assignee"] + ["Labels"] * LABEL_COLUMNS)                  # repeated header = multi-value Labels
        epic_id = {}
        for i, m in enumerate(milestones, start=1):                                     # epics first (parents before children)
            epic_id[m] = i
            w.writerow([i, "Epic", "", EPICS.get(m, m), f"Milestone {m}", "", "", m, "", "", ""] + [""] * LABEL_COLUMNS)
        next_id = len(milestones) + 1
        for r in rows:
            m = clean(r.get("Milestone")) or "-"
            lab = (labels(r) + [""] * LABEL_COLUMNS)[:LABEL_COLUMNS]
            w.writerow([next_id, ISSUE_TYPE.get(clean(r.get("Track")), "Task"), epic_id[m], summary(r), description(r, titles),
                        PRIORITY.get(clean(r.get("Priority")), "Medium"), STATUS.get(clean(r.get("Status")), "To Do"), m, r["#"],
                        sprint_of(r), members.get(clean(r.get("Owner")), "")] + lab)
            next_id += 1

    with (OUT_DIR / "jira-links.csv").open("w", newline="", encoding="utf-8") as fh:     # for the API step (links are not CSV-importable)
        w = csv.writer(fh)
        w.writerow(["plan_id", "blocked_by_plan_id"])
        for r in rows:
            for dep in re.findall(r"#(\d+)", clean(r.get("Depends"))):
                w.writerow([r["#"], dep])

    in_sprint = sum(1 for r in rows if sprint_of(r))
    print(f"exported {len(rows)} rows: plan/backlog.csv, plan/jira-import.csv ({len(rows) + len(milestones)} work items, "
          f"{in_sprint} in the sprint), plan/jira-links.csv")
    if not members:
        print("note: plan/members.json not found - Assignee left empty (copy plan/members.example.json)")
    if len(rows) + len(milestones) > 1500:
        print("note: Jira recommends at most ~1,500 work items per import file; split before importing")


def dependency_findings(rows):
    """Plan audit 2026-10-28: a row must not depend on a row planned in a LATER milestone, and dependencies must not form a cycle."""
    def stage(r):
        m = re.fullmatch(r"M(\d)", clean(r.get("Milestone")))
        return int(m.group(1)) if m else None
    by_id = {r["#"]: r for r in rows}
    deps = {r["#"]: [d for d in (int(x) for x in re.findall(r"#(\d+)", clean(r.get("Depends")))) if d in by_id and d != r["#"]] for r in rows}
    out = []
    for pid, ds in deps.items():
        for d in ds:
            a, b = stage(by_id[pid]), stage(by_id[d])
            if a is not None and b is not None and b > a:
                out.append(f"#{pid}: planned in M{a} but depends on #{d} planned in M{b}")
    state = {}
    def visit(n, path):
        state[n] = 1
        for d in deps[n]:
            if state.get(d) == 1:
                out.append("dependency cycle: " + " -> ".join(f"#{x}" for x in path + [n, d]))
            elif d not in state:
                visit(d, path + [n])
        state[n] = 2
    for n in deps:
        if n not in state:
            visit(n, [])
    return out


def check():
    _, rows = load_rows()
    problems = []
    ids = Counter(r["#"] for r in rows)
    problems += [f"duplicate Plan ID {i}" for i, n in ids.items() if n > 1]
    known = set(ids)
    for r in rows:
        pid = r["#"]
        status, evidence = clean(r.get("Status")), clean(r.get("Evidence"))
        # "Existing capabilities" rows are estimated from code (a class name is their evidence) and are listed for QA via the Test Plan
        if status == "Done" and len(evidence) < 20 and clean(r.get("Group")) != "Existing capabilities":
            problems.append(f"#{pid}: Done without evidence")
        if status not in STATUS:
            problems.append(f"#{pid}: unknown status '{status}'")
        if clean(r.get("Priority")) not in PRIORITY:
            problems.append(f"#{pid}: unknown priority '{clean(r.get('Priority'))}'")
        for dep in re.findall(r"#(\d+)", clean(r.get("Depends"))):
            if int(dep) not in known:
                problems.append(f"#{pid}: depends on unknown #{dep}")
        if clean(r.get("Sprint")):                                  # a planned row must be unambiguous
            if not clean(r.get("Owner")):
                problems.append(f"#{pid}: planned in {clean(r.get('Sprint'))} without an Owner")
            if len(clean(r.get("Prompt"))) < 60 or not clean(r.get("Requirement")).lower().startswith("done when"):
                problems.append(f"#{pid}: planned in a sprint but 'What to do' is short or 'Done when' is missing")
    problems += dependency_findings(rows)
    for p in problems:
        print("FINDING:", p)
    print(f"{len(rows)} rows checked, {len(problems)} finding(s)")
    return 1 if problems else 0


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "export":
        export()
    elif cmd == "check":
        sys.exit(check())
    else:
        sys.exit(__doc__)
